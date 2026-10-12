import asyncio
import tempfile
import unittest
from pathlib import Path

from docx import Document
from PIL import Image
from pypdf import PdfReader, PdfWriter

from herta_agent.workspace.paths import WorkspaceError
from herta_agent.workspace.service import WorkspaceService


class DocumentTests(unittest.IsolatedAsyncioTestCase):
    def test_artifact_quota_is_reserved_atomically_by_concurrent_publishers(self):
        from concurrent.futures import ThreadPoolExecutor
        from unittest.mock import patch

        with patch("herta_agent.documents.service.MAX_PENDING_BYTES", 10, create=True):
            self.d._publish(self.id, "base.md", b"123456", [])

            def publish(index):
                try:
                    self.d._publish(self.id, f"next-{index}.md", b"1234", [])
                    return True
                except WorkspaceError:
                    return False

            with ThreadPoolExecutor(2) as pool:
                results = list(pool.map(publish, range(2)))
            self.assertEqual(sum(results), 1)
            self.assertEqual(
                self.w._rows("SELECT SUM(size) total FROM document_artifacts")[0]["total"], 10
            )

    def test_failed_plan_save_does_not_leak_artifact(self):
        from unittest.mock import patch

        with patch.object(self.w.plans, "_save", side_effect=WorkspaceError("save failed")):
            with self.assertRaises(WorkspaceError):
                self.d._publish(self.id, "failed.md", b"generated", [])
        self.assertEqual(
            self.w._rows("SELECT COUNT(*) total FROM document_artifacts")[0]["total"], 0
        )

    def test_expired_artifacts_reclaim_quota_and_preserve_plan_metadata(self):
        from datetime import UTC, datetime, timedelta

        plan = self.d._publish(self.id, "expired.md", b"old", [])
        artifact = plan["operations"][-1]["artifact_id"]
        self.w._write(
            "UPDATE document_artifacts SET size=? WHERE id=?", (512 * 1024 * 1024, artifact)
        )
        self.w._write(
            "UPDATE file_plans SET created_at=? WHERE id=?",
            ((datetime.now(UTC) - timedelta(minutes=11)).isoformat(), plan["id"]),
        )
        next_plan = self.d._publish(self.id, "new.md", b"new", [])
        self.assertEqual(self.preserved_plan(plan["id"])["digest"], plan["digest"])
        old = self.w._rows("SELECT * FROM document_artifacts WHERE id=?", (artifact,))[0]
        self.assertIsNone(old["data"])
        self.assertEqual(old["size"], 0)
        self.assertEqual(next_plan["operations"][-1]["size"], 3)
        self.assertFalse((self.root / "expired.md").exists())

    def preserved_plan(self, identifier):
        return self.w.plans.get(identifier)

    def test_quota_reclamation_keeps_artifact_for_interrupted_intent(self):
        from datetime import UTC, datetime, timedelta

        plan = self.d._publish(self.id, "uncertain.md", b"owned", [])
        artifact = plan["operations"][-1]["artifact_id"]
        old = (datetime.now(UTC) - timedelta(minutes=11)).isoformat()
        self.w._write(
            "UPDATE document_artifacts SET size=? WHERE id=?", (512 * 1024 * 1024, artifact)
        )
        self.w._write("UPDATE file_plans SET created_at=? WHERE id=?", (old, plan["id"]))
        self.w._write(
            "INSERT INTO file_jobs VALUES(?,?,?,?,?,?,?)",
            ("uncertain", plan["id"], self.id, "uncertain", "interrupted", old, None),
        )
        self.w._write("INSERT INTO file_receipts VALUES(?,?,?,?)", ("uncertain", 0, "intent", "{}"))
        with self.assertRaises(WorkspaceError):
            self.d._publish(self.id, "new.md", b"new", [])
        self.assertEqual(
            self.w._rows("SELECT data FROM document_artifacts WHERE id=?", (artifact,))[0]["data"],
            b"owned",
        )

    async def test_semantic_fragments_preserve_tail_within_local_encoder_input_limit(self):
        from types import SimpleNamespace
        from unittest.mock import patch

        text = "量子" * 400 + "末尾关键知识"
        (self.root / "long.txt").write_text(text, encoding="utf-8")
        inputs = []
        self.d.embedding = SimpleNamespace(fingerprint="fixture")

        def encode(values):
            inputs.extend(values)
            return [[1.0, 0.0] for _ in values]

        with patch.object(self.d, "_vectors", side_effect=encode):
            await self.w.io(self.d.index, self.id, ["long.txt"], semantic=True)
        self.assertEqual("".join(inputs), text)
        self.assertLessEqual(max(map(len, inputs)), 300, "BGE的512-token窗口不能静默丢掉片段末尾")

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="Herta 文档 ")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name) / "资料"
        self.root.mkdir()
        self.w = WorkspaceService(Path(self.temp.name) / "workspace.sqlite3")
        self.addCleanup(self.w.close)
        self.id = self.w.grant(str(self.root))["id"]
        self.assertIsNotNone(getattr(self.w, "documents", None), "文档工作区尚未实现")
        self.d = self.w.documents
        doc = Document()
        doc.add_paragraph("量子计算学习笔记")
        doc.add_paragraph("量子比特可以描述叠加态。")
        doc.save(self.root / "材料.docx")
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=300)
        writer.add_blank_page(width=200, height=300)
        writer.write(self.root / "two.pdf")

    async def asyncTearDown(self):
        await self.w.shutdown()

    async def test_docx_paragraph_references_and_scanned_pdf_warning(self):
        result = await self.w.io(self.d.read, self.id, "材料.docx")
        self.assertIn("量子计算", result["segments"][0]["text"])
        self.assertEqual(result["segments"][1]["citation"], "材料.docx#段落2")
        pdf = await self.w.io(self.d.read, self.id, "two.pdf")
        self.assertEqual(pdf["units"], 2)
        self.assertTrue(pdf["warnings"])

    async def test_docx_tables_preserve_body_order_and_distinct_row_citations(self):
        doc = Document()
        doc.add_paragraph("前段")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "概念"
        table.cell(0, 1).text = "解释"
        doc.add_paragraph("后段")
        doc.save(self.root / "table.docx")
        result = await self.w.io(self.d.read, self.id, "table.docx")
        self.assertEqual(
            [s["citation"] for s in result["segments"]],
            ["table.docx#段落1", "table.docx#表1行1", "table.docx#段落2"],
        )

    async def test_document_bounds_corruption_links_and_shared_scope(self):
        (self.root / "broken.pdf").write_bytes(b"broken")
        with self.assertRaises(WorkspaceError):
            await self.w.io(self.d.read, self.id, "broken.pdf")
        (self.root / "big.txt").write_text("笔记" * 110000, encoding="utf-8")
        big = await self.w.io(self.d.read, self.id, "big.txt")
        self.assertTrue(big["truncated"])
        self.assertLessEqual(sum(len(s["text"]) for s in big["segments"]), 200000)
        for path in ("../out.txt", ".env", "材料.docx"):
            with self.assertRaises(WorkspaceError):
                await self.w.io(self.d.read, self.id, path, shared=True)

    async def test_pdf_transform_and_notes_only_publish_after_confirmation(self):
        plan = await self.w.io(
            self.d.transform,
            self.id,
            ["two.pdf"],
            "split",
            "outputs/one.pdf",
            pages=[2],
            rotation=0,
        )
        self.assertFalse((self.root / "outputs").exists())
        self.w.files.confirm(plan["id"], plan["digest"], "pdf")
        await self.w.files.wait()
        output = PdfReader(self.root / "outputs/one.pdf")
        self.assertEqual(len(output.pages), 1)
        self.assertEqual(self.w.files.history(self.id)[0]["status"], "completed")
        undo = self.w.plans.preview_undo(self.w.files.history(self.id)[0]["id"])
        self.assertFalse(undo["conflicts"])
        self.w.files.confirm(undo["id"], undo["digest"], "undo-pdf")
        await self.w.files.wait()
        self.assertFalse((self.root / "outputs").exists())
        self.d.bind_vault(self.id)
        note = await self.w.io(
            self.d.note,
            self.id,
            "notes/量子学习.md",
            "# 量子学习\n\n来源：材料.docx#段落2\n\n量子比特描述叠加态。",
        )
        self.assertFalse((self.root / "notes").exists())
        self.w.files.confirm(note["id"], note["digest"], "note")
        await self.w.files.wait()
        self.assertIn("材料.docx#段落2", (self.root / "notes/量子学习.md").read_text("utf-8"))

    async def test_generated_output_refuses_changed_source_and_existing_target(self):
        plan = await self.w.io(
            self.d.transform, self.id, ["two.pdf"], "rotate", "rotated.pdf", pages=[], rotation=90
        )
        (self.root / "two.pdf").write_bytes(b"later")
        self.w.files.confirm(plan["id"], plan["digest"], "changed-doc")
        await self.w.files.wait()
        self.assertEqual(self.w.files.history(self.id)[0]["status"], "failed")
        self.assertFalse((self.root / "rotated.pdf").exists())
        (self.root / "keep.md").write_text("keep")
        self.d.bind_vault(self.id)
        conflict = await self.w.io(self.d.note, self.id, "keep.md", "replacement")
        self.assertTrue(conflict["conflicts"])
        self.assertEqual((self.root / "keep.md").read_text(), "keep")

    async def test_image_and_merge_outputs_use_real_pdf_bytes(self):
        Image.new("RGB", (100, 150), "white").save(self.root / "image.png")
        for paths, kind, target, count in [
            (["image.png"], "images", "image.pdf", 1),
            (["two.pdf", "two.pdf"], "merge", "merged.pdf", 4),
        ]:
            plan = await self.w.io(
                self.d.transform, self.id, paths, kind, target, pages=[], rotation=0
            )
            self.w.files.confirm(plan["id"], plan["digest"], target)
            await self.w.files.wait()
            self.assertEqual(len(PdfReader(self.root / target).pages), count)

    async def test_personal_index_is_local_persistent_with_citations_and_revocation(self):
        indexed = await self.w.io(self.d.index, self.id, ["材料.docx"], semantic=False)
        self.assertEqual(indexed["documents"], 1)
        result = await self.w.io(self.d.search, self.id, "量子", semantic=False)
        self.assertTrue(result["results"])
        self.assertIn("#段落", result["results"][0]["citation"])
        with self.assertRaises(WorkspaceError):
            await self.w.io(self.d.search, self.id, "量子", shared=True, semantic=False)
        self.w.share(self.id, True)
        self.assertTrue(
            (await self.w.io(self.d.search, self.id, "量子", shared=True, semantic=False))[
                "results"
            ]
        )
        self.w.revoke(self.id)
        with self.assertRaises(WorkspaceError):
            await self.w.io(self.d.search, self.id, "量子", semantic=False)

    async def test_owned_parser_can_be_cancelled_and_closed(self):
        task = asyncio.create_task(self.w.io(self.d.read, self.id, "材料.docx"))
        await asyncio.sleep(0)
        task.cancel()
        with self.assertRaises(asyncio.CancelledError):
            await task
        self.assertEqual(self.w.active_count, 0)

    async def test_docx_pdf_uses_local_layout_engine(self):
        plan = await self.w.io(self.d.transform, self.id, ["材料.docx"], "docx_pdf", "材料.pdf")
        self.w.files.confirm(plan["id"], plan["digest"], "docx-layout")
        await self.w.files.wait()
        reader = PdfReader(self.root / "材料.pdf")
        self.assertEqual(len(reader.pages), 1)
        self.assertIn("量子", reader.pages[0].extract_text())

    async def test_published_artifact_releases_staged_bytes_but_retains_undo(self):
        self.d.bind_vault(self.id)
        plan = await self.w.io(self.d.note, self.id, "note.md", "# 本机笔记")
        self.w.files.confirm(plan["id"], plan["digest"], "release-artifact")
        await self.w.files.wait()
        self.assertEqual(
            self.w._rows("SELECT SUM(size) AS total FROM document_artifacts")[0]["total"], 0
        )
        self.assertFalse(
            self.w.plans.preview_undo(self.w.files.history(self.id)[0]["id"])["conflicts"]
        )

    async def test_index_catalog_removal_and_modified_sources_are_not_returned(self):
        await self.w.io(self.d.index, self.id, ["材料.docx"], semantic=False)
        self.assertEqual(self.d.catalog(self.id)[0]["path"], "材料.docx")
        (self.root / "材料.docx").write_bytes(b"new user bytes")
        self.assertFalse(
            (await self.w.io(self.d.search, self.id, "量子", semantic=False))["results"]
        )
        self.d.forget(self.id, "材料.docx")
        self.assertEqual(self.d.catalog(self.id), [])

    async def test_cancel_parser_joins_its_actual_child_process(self):
        import subprocess
        import threading
        from unittest.mock import patch

        entered, release = threading.Event(), threading.Event()
        children = []
        original = subprocess.Popen

        def starting(*args, **kwargs):
            child = original(*args, **kwargs)
            children.append(child)
            entered.set()
            release.wait(3)
            return child

        with patch("herta_agent.documents.process.subprocess.Popen", starting):
            task = asyncio.create_task(self.w.io(self.d.read, self.id, "材料.docx"))
            self.assertTrue(await asyncio.to_thread(entered.wait, 2))
            task.cancel()
            release.set()
            with self.assertRaises(asyncio.CancelledError):
                await task
        self.assertIsNotNone(children[0].poll())
        self.assertEqual(self.w.active_count, 0)
